#!/usr/bin/env python3
"""CogniBias — régénère data/biases.json et data/data.js à partir de Biais101.xlsx.

Usage :  python3 tools/excel_to_json.py chemin/vers/Biais101.xlsx
Dépendance : openpyxl (pip install openpyxl)

Le classeur doit contenir les onglets « Biais » (FR), « Biases (EN) », « Processus » et « Codex Benson ».
Les deux onglets de biais sont reliés par la colonne ID.
"""
import sys, re, json, unicodedata
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText

ROB = {"Répliqué": "replicated", "Replicated": "replicated", "Débattu": "debated",
       "Debated": "debated", "Fragile": "fragile"}

def slug(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")

def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def to_html(v):
    """Texte enrichi Excel -> HTML (l'italique est conservé), sans le DOI final."""
    if v is None:
        return ""
    if isinstance(v, CellRichText):
        out = []
        for part in v:
            if isinstance(part, str):
                out.append(esc(part))
            else:
                t = esc(part.text)
                out.append(f"<i>{t}</i>" if part.font is not None and part.font.i else t)
        html = "".join(out)
    else:
        html = esc(str(v))
    return re.sub(r"\s*https://doi\.org/\S+$", "", html).strip()

def rows(ws):
    it = ws.iter_rows()
    head = [c.value for c in next(it)]
    for r in it:
        if r[0].value:
            yield {h: c.value for h, c in zip(head, r)}

def main(xlsx, out_dir):
    wb = load_workbook(xlsx, rich_text=True)
    fr = {r["ID"]: r for r in rows(wb["Biais"])}
    en = {r["ID"]: r for r in rows(wb["Biases (EN)"])}
    missing = set(fr) ^ set(en)
    if missing:
        sys.exit(f"Fiches présentes dans une seule langue : {sorted(missing)}")

    processes = [{"id": r["Code"],
                  "fr": {"name": r["Processus"], "scope": r["Ce qu'il recouvre"]},
                  "en": {"name": r["Process (EN)"], "scope": r["Scope (EN)"]}}
                 for r in rows(wb["Processus"]) if r.get("Code")]

    problems, strategies, seen = [], [], set()
    for r in rows(wb["Codex Benson"]):
        code = str(r["Code"]).strip()
        if re.fullmatch(r"C\d[a-z]", code):
            strategies.append({"id": code, "fr": r["Stratégie"], "en": r["Strategy (EN)"]})
            if code[:2] not in seen:
                seen.add(code[:2])
                problems.append({"id": code[:2], "fr": r["Problème"], "en": r["Problem (EN)"]})

    biases = []
    for bid in sorted(fr):
        f, e = fr[bid], en[bid]
        name_en = re.match(r"^(.*?)(?:\s*\((.*)\))?$", e["Name (EN)"])
        rel = [x.strip() for x in str(f["Voisins (ID)"] or "").split(";") if x.strip()]
        biases.append({
            "id": bid, "n": int(f["N°"]), "process": f["Code"], "robustness": ROB[f["Robustesse"]],
            "benson": {"problem": str(f["Stratégie (Benson)"])[:2], "strategy": str(f["Stratégie (Benson)"]).split(" ·")[0],
                       "listed": f["Rattachement Benson"] == "Cité par Benson"},
            "doi": f["DOI"] or None, "reference": to_html(f["Étude fondatrice (APA 7)"]), "related": rel,
            "fr": {"name": f["Nom (FR)"], "slug": slug(f["Nom (FR)"]), "definition": f["Définition"],
                   "example": f["Exemple du quotidien"], "mechanism": f["Mécanisme / heuristique"],
                   "robustness_note": f["Note de robustesse"], "countermeasures": f["Parades"]},
            "en": {"name": name_en.group(1), "aka": name_en.group(2) or "", "slug": slug(name_en.group(1)),
                   "definition": e["Definition"], "example": e["Everyday example"],
                   "mechanism": e["Mechanism / heuristic"], "robustness_note": e["Robustness note"],
                   "countermeasures": e["Countermeasures"]},
        })

    ids = {b["id"] for b in biases}
    broken = [(b["id"], r) for b in biases for r in b["related"] if r not in ids]
    if broken:
        sys.exit(f"Liens vers des fiches inexistantes : {broken}")

    data = {"processes": processes, "benson": {"problems": problems, "strategies": strategies}, "biases": biases}
    out = Path(out_dir); out.mkdir(parents=True, exist_ok=True)
    js = json.dumps(data, ensure_ascii=False, indent=1)
    (out / "biases.json").write_text(js, encoding="utf-8")
    (out / "data.js").write_text("// Généré par tools/excel_to_json.py — ne pas modifier à la main\nwindow.COGNIBIAS = " + js + ";\n", encoding="utf-8")
    print(f"{len(biases)} fiches, {len(processes)} processus -> {out}/")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else Path(__file__).resolve().parent.parent / "data")
